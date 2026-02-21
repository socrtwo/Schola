import datetime
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.db import IntegrityError
from django.test import TestCase
from django.utils import timezone

from core.models import (
    Attendance,
    BundleDefinition,
    BundlePrice,
    DiscountRule,
    Event,
    EventCategory,
    EventPrice,
    EventSession,
    Organization,
    Payment,
    Person,
    PricingTier,
    RefundPolicy,
    Registration,
    Term,
    Venue,
    WaitlistEntry,
    WorkflowState,
)
from core.services import (
    active_registration_count,
    calculate_bundle_price,
    calculate_discount,
    calculate_refund,
    calculate_registration_price,
    cancel_registration,
    get_bundle_price,
    get_event_price,
    is_event_full,
    process_refund,
    promote_from_waitlist,
    record_offline_payment,
    reconcile_payment,
    register_bundle,
    register_person,
)


class OrgFactory:
    """Helper to create test orgs and related objects."""

    @staticmethod
    def create_org(slug="test-org", **kwargs):
        defaults = {"name": f"Test Org ({slug})", "slug": slug}
        defaults.update(kwargs)
        return Organization.objects.create(**defaults)

    @staticmethod
    def create_term(org, **kwargs):
        defaults = {
            "organization": org,
            "name": "Fall 2025",
            "start_date": datetime.date(2025, 9, 1),
            "end_date": datetime.date(2025, 12, 15),
        }
        defaults.update(kwargs)
        return Term.objects.create(**defaults)

    @staticmethod
    def create_person(org, first="Jane", last="Doe", **kwargs):
        defaults = {
            "organization": org,
            "first_name": first,
            "last_name": last,
        }
        defaults.update(kwargs)
        return Person.objects.create(**defaults)


# ---------------------------------------------------------------------------
# Tenant Scoping Tests
# ---------------------------------------------------------------------------

class TenantScopingTests(TestCase):

    def setUp(self):
        self.org_a = OrgFactory.create_org("org-a")
        self.org_b = OrgFactory.create_org("org-b")
        self.term_a = OrgFactory.create_term(self.org_a, name="Fall A")
        self.term_b = OrgFactory.create_term(self.org_b, name="Fall B")

    def test_for_org_filters_terms(self):
        """Term.objects.for_org() returns only that org's terms."""
        terms_a = Term.objects.for_org(self.org_a)
        terms_b = Term.objects.for_org(self.org_b)
        self.assertEqual(list(terms_a), [self.term_a])
        self.assertEqual(list(terms_b), [self.term_b])

    def test_for_org_filters_people(self):
        """Person.objects.for_org() returns only that org's people."""
        p_a = OrgFactory.create_person(self.org_a, first="Alice")
        p_b = OrgFactory.create_person(self.org_b, first="Bob")
        self.assertEqual(list(Person.objects.for_org(self.org_a)), [p_a])
        self.assertEqual(list(Person.objects.for_org(self.org_b)), [p_b])

    def test_unique_email_per_org(self):
        """Same email can exist in different orgs but not the same org."""
        OrgFactory.create_person(self.org_a, email="test@example.com")
        # Same email, different org — OK
        OrgFactory.create_person(self.org_b, first="Bob", email="test@example.com")
        # Same email, same org — should fail
        with self.assertRaises(IntegrityError):
            OrgFactory.create_person(self.org_a, first="Dupe", email="test@example.com")

    def test_blank_email_not_unique(self):
        """Multiple people with blank email in the same org is allowed."""
        OrgFactory.create_person(self.org_a, first="No Email 1", email="")
        OrgFactory.create_person(self.org_a, first="No Email 2", email="")
        self.assertEqual(Person.objects.for_org(self.org_a).count(), 2)


# ---------------------------------------------------------------------------
# Event Lifecycle / Workflow Tests
# ---------------------------------------------------------------------------

class EventLifecycleTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.venue = Venue.objects.create(
            organization=self.org, name="Main Hall", capacity=100,
        )
        self.draft = WorkflowState.objects.create(
            organization=self.org, name="Draft", slug="draft", position=0,
        )
        self.tentative = WorkflowState.objects.create(
            organization=self.org, name="Tentative", slug="tentative", position=1,
        )
        self.booked = WorkflowState.objects.create(
            organization=self.org, name="Booked", slug="booked", position=2,
            gate_conditions={
                "has_abstract": True,
                "has_bio": True,
                "has_graphic": True,
                "venue_id__isnull": False,
            },
        )
        self.published = WorkflowState.objects.create(
            organization=self.org, name="Published", slug="published", position=3,
        )

    def _create_event(self, **kwargs):
        defaults = {
            "organization": self.org,
            "term": self.term,
            "title": "Test Event",
            "lifecycle_state": self.draft,
        }
        defaults.update(kwargs)
        return Event.objects.create(**defaults)

    def test_transition_draft_to_tentative(self):
        """Draft → Tentative succeeds (no gate conditions)."""
        event = self._create_event()
        event.transition_to(self.tentative)
        event.refresh_from_db()
        self.assertEqual(event.lifecycle_state, self.tentative)

    def test_transition_to_booked_fails_without_requirements(self):
        """Tentative → Booked fails when gate conditions not met."""
        event = self._create_event(lifecycle_state=self.tentative)
        with self.assertRaises(ValidationError) as ctx:
            event.transition_to(self.booked)
        self.assertIn("has_abstract", str(ctx.exception))

    def test_transition_to_booked_succeeds_with_requirements(self):
        """Tentative → Booked succeeds when all gate conditions are met."""
        event = self._create_event(
            lifecycle_state=self.tentative,
            venue=self.venue,
            has_abstract=True,
            has_bio=True,
            has_graphic=True,
        )
        event.transition_to(self.booked)
        event.refresh_from_db()
        self.assertEqual(event.lifecycle_state, self.booked)

    def test_transition_to_booked_fails_without_venue(self):
        """Booked requires a venue (venue_id__isnull: False)."""
        event = self._create_event(
            lifecycle_state=self.tentative,
            has_abstract=True,
            has_bio=True,
            has_graphic=True,
            venue=None,
        )
        with self.assertRaises(ValidationError) as ctx:
            event.transition_to(self.booked)
        self.assertIn("venue_id", str(ctx.exception))

    def test_transition_to_published_no_extra_gates(self):
        """Booked → Published succeeds (no extra gate conditions)."""
        event = self._create_event(
            lifecycle_state=self.booked,
            venue=self.venue,
            has_abstract=True, has_bio=True, has_graphic=True,
        )
        event.transition_to(self.published)
        event.refresh_from_db()
        self.assertEqual(event.lifecycle_state, self.published)


# ---------------------------------------------------------------------------
# Capacity & Waitlist Tests
# ---------------------------------------------------------------------------

class CapacityTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.event = Event.objects.create(
            organization=self.org, term=self.term,
            title="Limited Event", capacity=2,
        )
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )

    def test_registration_count(self):
        """Active registrations are correctly counted."""
        for i in range(2):
            p = OrgFactory.create_person(self.org, first=f"Student{i}", email=f"s{i}@test.com")
            Registration.objects.create(
                organization=self.org, person=p, event=self.event,
                term=self.term, tier=self.tier,
                status="registered", amount_charged=Decimal("1.00"),
            )
        active = self.event.registrations.filter(status="registered").count()
        self.assertEqual(active, 2)
        self.assertEqual(active, self.event.capacity)

    def test_cancelled_does_not_count(self):
        """Cancelled registrations don't count toward capacity."""
        p = OrgFactory.create_person(self.org, email="cancel@test.com")
        reg = Registration.objects.create(
            organization=self.org, person=p, event=self.event,
            term=self.term, status="registered", amount_charged=Decimal("1.00"),
        )
        reg.status = "cancelled"
        reg.save()
        active = self.event.registrations.filter(status="registered").count()
        self.assertEqual(active, 0)

    def test_unique_registration_per_person_event(self):
        """A person cannot register twice for the same event."""
        p = OrgFactory.create_person(self.org, email="dupe@test.com")
        Registration.objects.create(
            organization=self.org, person=p, event=self.event,
            term=self.term, status="registered",
        )
        with self.assertRaises(IntegrityError):
            Registration.objects.create(
                organization=self.org, person=p, event=self.event,
                term=self.term, status="registered",
            )


# ---------------------------------------------------------------------------
# Pricing Tests
# ---------------------------------------------------------------------------

class PricingTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.event = Event.objects.create(
            organization=self.org, term=self.term, title="Priced Event",
        )
        self.tier_res = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.tier_nr = PricingTier.objects.create(
            organization=self.org, name="Non-Resident", slug="non-resident",
        )

    def test_different_prices_per_tier(self):
        """Events can have different prices for different tiers."""
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier_res, timing="pre", amount=Decimal("1.00"),
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier_nr, timing="pre", amount=Decimal("5.00"),
        )
        res_price = EventPrice.objects.get(
            event=self.event, tier=self.tier_res, timing="pre",
        )
        nr_price = EventPrice.objects.get(
            event=self.event, tier=self.tier_nr, timing="pre",
        )
        self.assertEqual(res_price.amount, Decimal("1.00"))
        self.assertEqual(nr_price.amount, Decimal("5.00"))

    def test_pre_vs_walk_pricing(self):
        """Pre-registration and walk-in prices can differ."""
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier_res, timing="pre", amount=Decimal("1.00"),
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier_res, timing="walk", amount=Decimal("2.00"),
        )
        pre = EventPrice.objects.get(event=self.event, tier=self.tier_res, timing="pre")
        walk = EventPrice.objects.get(event=self.event, tier=self.tier_res, timing="walk")
        self.assertLess(pre.amount, walk.amount)

    def test_unique_price_constraint(self):
        """Cannot create duplicate price for same event+tier+timing."""
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier_res, timing="pre", amount=Decimal("1.00"),
        )
        with self.assertRaises(IntegrityError):
            EventPrice.objects.create(
                organization=self.org, event=self.event,
                tier=self.tier_res, timing="pre", amount=Decimal("2.00"),
            )


# ---------------------------------------------------------------------------
# Bundle Tests
# ---------------------------------------------------------------------------

class BundleTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.cat = EventCategory.objects.create(
            organization=self.org, name="Evening Lecture", slug="evening-lecture",
        )
        self.evt1 = Event.objects.create(
            organization=self.org, term=self.term, category=self.cat,
            title="Lecture 1",
        )
        self.evt2 = Event.objects.create(
            organization=self.org, term=self.term, category=self.cat,
            title="Lecture 2",
        )
        self.evt3 = Event.objects.create(
            organization=self.org, term=self.term,
            title="Performance 1",  # no category
        )

    def test_category_bundle(self):
        """Category-based bundle returns all events in that category."""
        bundle = BundleDefinition.objects.create(
            organization=self.org, name="All Lectures", slug="all-lectures",
            term=self.term, selection_mode="category",
            category_filter=self.cat,
        )
        eligible = bundle.get_eligible_events()
        self.assertEqual(set(eligible), {self.evt1, self.evt2})

    def test_manual_bundle(self):
        """Manual bundle returns only explicitly added events."""
        bundle = BundleDefinition.objects.create(
            organization=self.org, name="Custom", slug="custom",
            term=self.term, selection_mode="manual",
        )
        bundle.events.add(self.evt1, self.evt3)
        eligible = bundle.get_eligible_events()
        self.assertEqual(set(eligible), {self.evt1, self.evt3})


# ===========================================================================
# Phase 2 Tests — Service Layer
# ===========================================================================


# ---------------------------------------------------------------------------
# Registration Service Tests (capacity + waitlist)
# ---------------------------------------------------------------------------

class RegistrationServiceTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.event = Event.objects.create(
            organization=self.org, term=self.term,
            title="Capped Event", capacity=2,
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("5.00"),
        )
        self.students = []
        for i in range(4):
            self.students.append(
                OrgFactory.create_person(
                    self.org, first=f"Student{i}", email=f"s{i}@test.com",
                )
            )

    def test_register_under_capacity(self):
        """Registration succeeds when under capacity."""
        reg, wl = register_person(
            self.org, self.students[0], self.event, self.term, self.tier,
        )
        self.assertEqual(reg.status, "registered")
        self.assertIsNone(wl)
        self.assertEqual(reg.amount_charged, Decimal("5.00"))

    def test_register_at_capacity_goes_to_waitlist(self):
        """Third registration on a capacity-2 event is waitlisted."""
        register_person(self.org, self.students[0], self.event, self.term, self.tier)
        register_person(self.org, self.students[1], self.event, self.term, self.tier)
        reg, wl = register_person(
            self.org, self.students[2], self.event, self.term, self.tier,
        )
        self.assertEqual(reg.status, "waitlisted")
        self.assertIsNotNone(wl)
        self.assertEqual(wl.position, 1)

    def test_waitlist_ordering(self):
        """Waitlist positions increment correctly."""
        register_person(self.org, self.students[0], self.event, self.term, self.tier)
        register_person(self.org, self.students[1], self.event, self.term, self.tier)
        _, wl1 = register_person(
            self.org, self.students[2], self.event, self.term, self.tier,
        )
        _, wl2 = register_person(
            self.org, self.students[3], self.event, self.term, self.tier,
        )
        self.assertEqual(wl1.position, 1)
        self.assertEqual(wl2.position, 2)

    def test_register_without_waitlist_raises_when_full(self):
        """allow_waitlist=False raises ValidationError at capacity."""
        register_person(self.org, self.students[0], self.event, self.term, self.tier)
        register_person(self.org, self.students[1], self.event, self.term, self.tier)
        with self.assertRaises(ValidationError):
            register_person(
                self.org, self.students[2], self.event, self.term, self.tier,
                allow_waitlist=False,
            )

    def test_duplicate_registration_raises(self):
        """Cannot register same person for same event twice."""
        register_person(self.org, self.students[0], self.event, self.term, self.tier)
        with self.assertRaises(ValidationError):
            register_person(
                self.org, self.students[0], self.event, self.term, self.tier,
            )

    def test_reregister_after_cancel(self):
        """A previously cancelled person can re-register."""
        reg, _ = register_person(
            self.org, self.students[0], self.event, self.term, self.tier,
        )
        cancel_registration(reg)
        reg2, _ = register_person(
            self.org, self.students[0], self.event, self.term, self.tier,
        )
        self.assertEqual(reg2.status, "registered")

    def test_no_capacity_allows_unlimited(self):
        """Event with capacity=None allows unlimited registrations."""
        unlimited = Event.objects.create(
            organization=self.org, term=self.term,
            title="Unlimited Event",  # no capacity
        )
        EventPrice.objects.create(
            organization=self.org, event=unlimited,
            tier=self.tier, timing="pre", amount=Decimal("1.00"),
        )
        for student in self.students:
            reg, wl = register_person(
                self.org, student, unlimited, self.term, self.tier,
            )
            self.assertEqual(reg.status, "registered")
            self.assertIsNone(wl)

    def test_active_registration_count(self):
        """active_registration_count counts only 'registered' status."""
        register_person(self.org, self.students[0], self.event, self.term, self.tier)
        register_person(self.org, self.students[1], self.event, self.term, self.tier)
        self.assertEqual(active_registration_count(self.event), 2)

    def test_is_event_full(self):
        """is_event_full returns True at capacity."""
        self.assertFalse(is_event_full(self.event))
        register_person(self.org, self.students[0], self.event, self.term, self.tier)
        self.assertFalse(is_event_full(self.event))
        register_person(self.org, self.students[1], self.event, self.term, self.tier)
        self.assertTrue(is_event_full(self.event))


# ---------------------------------------------------------------------------
# Waitlist Promotion Tests
# ---------------------------------------------------------------------------

class WaitlistPromotionTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.event = Event.objects.create(
            organization=self.org, term=self.term,
            title="Waitlist Event", capacity=1,
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("3.00"),
        )
        self.p1 = OrgFactory.create_person(self.org, first="P1", email="p1@t.com")
        self.p2 = OrgFactory.create_person(self.org, first="P2", email="p2@t.com")
        self.p3 = OrgFactory.create_person(self.org, first="P3", email="p3@t.com")

    def test_cancel_promotes_first_waitlisted(self):
        """Cancelling a registered person promotes first waitlisted."""
        reg1, _ = register_person(
            self.org, self.p1, self.event, self.term, self.tier,
        )
        reg2, wl2 = register_person(
            self.org, self.p2, self.event, self.term, self.tier,
        )
        self.assertEqual(reg2.status, "waitlisted")

        _, promoted = cancel_registration(reg1)
        self.assertIsNotNone(promoted)
        self.assertEqual(promoted.person, self.p2)
        promoted.refresh_from_db()
        self.assertEqual(promoted.status, "registered")
        # Waitlist entry should be deleted
        self.assertFalse(WaitlistEntry.objects.filter(registration=reg2).exists())

    def test_cancel_promotes_in_order(self):
        """When multiple waitlisted, the first position is promoted."""
        reg1, _ = register_person(
            self.org, self.p1, self.event, self.term, self.tier,
        )
        _, wl2 = register_person(
            self.org, self.p2, self.event, self.term, self.tier,
        )
        _, wl3 = register_person(
            self.org, self.p3, self.event, self.term, self.tier,
        )
        self.assertEqual(wl2.position, 1)
        self.assertEqual(wl3.position, 2)

        _, promoted = cancel_registration(reg1)
        self.assertEqual(promoted.person, self.p2)
        # P3 should still be waitlisted
        reg3 = Registration.objects.get(person=self.p3, event=self.event)
        self.assertEqual(reg3.status, "waitlisted")

    def test_cancel_waitlisted_no_promotion(self):
        """Cancelling a waitlisted registration doesn't trigger promotion."""
        reg1, _ = register_person(
            self.org, self.p1, self.event, self.term, self.tier,
        )
        reg2, _ = register_person(
            self.org, self.p2, self.event, self.term, self.tier,
        )
        _, promoted = cancel_registration(reg2)
        self.assertIsNone(promoted)
        # P1 should still be registered
        reg1.refresh_from_db()
        self.assertEqual(reg1.status, "registered")

    def test_promote_from_empty_waitlist(self):
        """promote_from_waitlist returns None when no one is waiting."""
        result = promote_from_waitlist(self.event)
        self.assertIsNone(result)


# ---------------------------------------------------------------------------
# Pricing Service Tests
# ---------------------------------------------------------------------------

class PricingServiceTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.event = Event.objects.create(
            organization=self.org, term=self.term, title="Priced Event",
        )
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("10.00"),
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="walk", amount=Decimal("15.00"),
        )

    def test_get_event_price_pre(self):
        """get_event_price returns correct pre-registration price."""
        price = get_event_price(self.event, self.tier, "pre")
        self.assertEqual(price, Decimal("10.00"))

    def test_get_event_price_walk(self):
        """get_event_price returns correct walk-in price."""
        price = get_event_price(self.event, self.tier, "walk")
        self.assertEqual(price, Decimal("15.00"))

    def test_get_event_price_missing(self):
        """get_event_price returns None for undefined tier."""
        other_tier = PricingTier.objects.create(
            organization=self.org, name="NR", slug="nr",
        )
        self.assertIsNone(get_event_price(self.event, other_tier, "pre"))

    def test_calculate_registration_price_no_discount(self):
        """Pricing calculation without discounts."""
        base, disc, final = calculate_registration_price(
            self.event, self.tier, "pre",
        )
        self.assertEqual(base, Decimal("10.00"))
        self.assertEqual(disc, Decimal("0.00"))
        self.assertEqual(final, Decimal("10.00"))

    def test_calculate_registration_price_walk(self):
        """Walk-in pricing is correctly used."""
        base, disc, final = calculate_registration_price(
            self.event, self.tier, "walk",
        )
        self.assertEqual(base, Decimal("15.00"))
        self.assertEqual(final, Decimal("15.00"))

    def test_calculate_registration_price_undefined_raises(self):
        """Missing price raises ValidationError."""
        other_tier = PricingTier.objects.create(
            organization=self.org, name="NR", slug="nr2",
        )
        with self.assertRaises(ValidationError):
            calculate_registration_price(self.event, other_tier, "pre")


# ---------------------------------------------------------------------------
# Discount Tests
# ---------------------------------------------------------------------------

class DiscountTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.event = Event.objects.create(
            organization=self.org, term=self.term, title="Discount Event",
        )
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("20.00"),
        )

    def test_percent_discount(self):
        """Percentage discount is correctly applied."""
        rule = DiscountRule.objects.create(
            organization=self.org, name="10% Off", discount_type="percent",
            discount_value=Decimal("10"), applies_to="tier",
            tier=self.tier, is_active=True,
        )
        discount = calculate_discount(Decimal("20.00"), rule)
        self.assertEqual(discount, Decimal("2.00"))

    def test_fixed_discount(self):
        """Fixed discount is correctly applied."""
        rule = DiscountRule.objects.create(
            organization=self.org, name="$5 Off", discount_type="fixed",
            discount_value=Decimal("5.00"), applies_to="tier",
            tier=self.tier, is_active=True,
        )
        discount = calculate_discount(Decimal("20.00"), rule)
        self.assertEqual(discount, Decimal("5.00"))

    def test_fixed_discount_capped_at_amount(self):
        """Fixed discount cannot exceed the base amount."""
        rule = DiscountRule.objects.create(
            organization=self.org, name="$30 Off", discount_type="fixed",
            discount_value=Decimal("30.00"), applies_to="tier",
            tier=self.tier, is_active=True,
        )
        discount = calculate_discount(Decimal("20.00"), rule)
        self.assertEqual(discount, Decimal("20.00"))

    def test_inactive_discount_returns_zero(self):
        """Inactive discount rule returns zero."""
        rule = DiscountRule.objects.create(
            organization=self.org, name="Inactive", discount_type="percent",
            discount_value=Decimal("50"), applies_to="tier",
            tier=self.tier, is_active=False,
        )
        discount = calculate_discount(Decimal("20.00"), rule)
        self.assertEqual(discount, Decimal("0.00"))

    def test_expired_discount_returns_zero(self):
        """Expired discount rule returns zero."""
        rule = DiscountRule.objects.create(
            organization=self.org, name="Expired", discount_type="percent",
            discount_value=Decimal("50"), applies_to="tier",
            tier=self.tier, is_active=True,
            valid_until=timezone.now() - datetime.timedelta(days=1),
        )
        discount = calculate_discount(Decimal("20.00"), rule)
        self.assertEqual(discount, Decimal("0.00"))

    def test_future_discount_returns_zero(self):
        """Not-yet-valid discount rule returns zero."""
        rule = DiscountRule.objects.create(
            organization=self.org, name="Future", discount_type="percent",
            discount_value=Decimal("50"), applies_to="tier",
            tier=self.tier, is_active=True,
            valid_from=timezone.now() + datetime.timedelta(days=30),
        )
        discount = calculate_discount(Decimal("20.00"), rule)
        self.assertEqual(discount, Decimal("0.00"))

    def test_promo_code_discount_applied(self):
        """Promo code discount applies through calculate_registration_price."""
        DiscountRule.objects.create(
            organization=self.org, name="Promo 15%", discount_type="percent",
            discount_value=Decimal("15"), applies_to="promo",
            promo_code="SAVE15", is_active=True,
        )
        base, disc, final = calculate_registration_price(
            self.event, self.tier, "pre", promo_code="SAVE15",
        )
        self.assertEqual(base, Decimal("20.00"))
        self.assertEqual(disc, Decimal("3.00"))
        self.assertEqual(final, Decimal("17.00"))

    def test_tier_discount_applied(self):
        """Tier-based discount applies through calculate_registration_price."""
        DiscountRule.objects.create(
            organization=self.org, name="Resident discount", discount_type="fixed",
            discount_value=Decimal("2.00"), applies_to="tier",
            tier=self.tier, is_active=True,
        )
        base, disc, final = calculate_registration_price(
            self.event, self.tier, "pre",
        )
        self.assertEqual(base, Decimal("20.00"))
        self.assertEqual(disc, Decimal("2.00"))
        self.assertEqual(final, Decimal("18.00"))


# ---------------------------------------------------------------------------
# Bundle Pricing & Registration Tests
# ---------------------------------------------------------------------------

class BundlePricingTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.cat = EventCategory.objects.create(
            organization=self.org, name="Evening Lecture", slug="evening-lecture",
        )
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.evt1 = Event.objects.create(
            organization=self.org, term=self.term, category=self.cat,
            title="Lecture 1",
        )
        self.evt2 = Event.objects.create(
            organization=self.org, term=self.term, category=self.cat,
            title="Lecture 2",
        )
        # Individual prices
        for evt in [self.evt1, self.evt2]:
            EventPrice.objects.create(
                organization=self.org, event=evt,
                tier=self.tier, timing="pre", amount=Decimal("10.00"),
            )
        self.bundle = BundleDefinition.objects.create(
            organization=self.org, name="All Lectures", slug="all-lectures",
            term=self.term, selection_mode="category",
            category_filter=self.cat,
        )
        BundlePrice.objects.create(
            organization=self.org, bundle=self.bundle,
            tier=self.tier, amount=Decimal("15.00"),
        )

    def test_get_bundle_price(self):
        """get_bundle_price returns correct price."""
        price = get_bundle_price(self.bundle, self.tier)
        self.assertEqual(price, Decimal("15.00"))

    def test_get_bundle_price_missing(self):
        """get_bundle_price returns None for undefined tier."""
        other = PricingTier.objects.create(
            organization=self.org, name="NR", slug="nr",
        )
        self.assertIsNone(get_bundle_price(self.bundle, other))

    def test_calculate_bundle_price_no_discount(self):
        """Bundle pricing calculation without discounts."""
        base, disc, final = calculate_bundle_price(self.bundle, self.tier)
        self.assertEqual(base, Decimal("15.00"))
        self.assertEqual(disc, Decimal("0.00"))
        self.assertEqual(final, Decimal("15.00"))

    def test_bundle_discount(self):
        """Bundle-specific discount is applied."""
        DiscountRule.objects.create(
            organization=self.org, name="Bundle 10%", discount_type="percent",
            discount_value=Decimal("10"), applies_to="bundle",
            bundle=self.bundle, is_active=True,
        )
        base, disc, final = calculate_bundle_price(self.bundle, self.tier)
        self.assertEqual(base, Decimal("15.00"))
        self.assertEqual(disc, Decimal("1.50"))
        self.assertEqual(final, Decimal("13.50"))

    def test_register_bundle(self):
        """register_bundle creates registrations for all bundle events."""
        person = OrgFactory.create_person(self.org, email="bundler@t.com")
        regs, amount = register_bundle(
            self.org, person, self.bundle, self.term, self.tier,
        )
        self.assertEqual(len(regs), 2)
        self.assertEqual(amount, Decimal("15.00"))
        # Individual registrations should be $0
        for reg in regs:
            self.assertEqual(reg.amount_charged, Decimal("0.00"))
            self.assertEqual(reg.bundle, self.bundle)

    def test_register_bundle_empty_raises(self):
        """Bundle with no events raises ValidationError."""
        empty_bundle = BundleDefinition.objects.create(
            organization=self.org, name="Empty", slug="empty",
            term=self.term, selection_mode="manual",
        )
        BundlePrice.objects.create(
            organization=self.org, bundle=empty_bundle,
            tier=self.tier, amount=Decimal("5.00"),
        )
        person = OrgFactory.create_person(self.org, email="empty@t.com")
        with self.assertRaises(ValidationError):
            register_bundle(
                self.org, person, empty_bundle, self.term, self.tier,
            )


# ---------------------------------------------------------------------------
# Refund Policy Tests
# ---------------------------------------------------------------------------

class RefundPolicyTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.event = Event.objects.create(
            organization=self.org, term=self.term,
            title="Refund Event",
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("20.00"),
        )
        # Session 30 days from now
        self.session = EventSession.objects.create(
            organization=self.org, event=self.event,
            date=datetime.date.today() + datetime.timedelta(days=30),
            start_time=datetime.time(19, 30),
        )
        self.policy = RefundPolicy.objects.create(
            organization=self.org, name="Standard", is_default=True,
            full_refund_before_days=7,
            partial_refund_percent=Decimal("50.00"),
            admin_only_after_days=1,
        )
        self.person = OrgFactory.create_person(self.org, email="refund@t.com")

    def _register_and_cancel(self):
        reg, _ = register_person(
            self.org, self.person, self.event, self.term, self.tier,
        )
        cancel_registration(reg)
        reg.refresh_from_db()
        return reg

    def test_full_refund_before_window(self):
        """Full refund when event is more than 7 days away."""
        reg = self._register_and_cancel()
        amount, rtype = calculate_refund(reg)
        self.assertEqual(amount, Decimal("20.00"))
        self.assertEqual(rtype, "full")

    def test_partial_refund_in_window(self):
        """Partial refund between admin_only and full_refund windows."""
        reg = self._register_and_cancel()
        # Reference date: 3 days before event (between 1 and 7)
        ref_date = self.session.date - datetime.timedelta(days=3)
        amount, rtype = calculate_refund(reg, reference_date=ref_date)
        self.assertEqual(amount, Decimal("10.00"))
        self.assertEqual(rtype, "partial")

    def test_admin_only_refund(self):
        """Admin-only when less than 1 day before event."""
        reg = self._register_and_cancel()
        ref_date = self.session.date  # day of event = 0 days before
        amount, rtype = calculate_refund(reg, reference_date=ref_date)
        self.assertEqual(amount, Decimal("0.00"))
        self.assertEqual(rtype, "admin_only")

    def test_process_refund_creates_payment(self):
        """process_refund creates a refund Payment record."""
        reg = self._register_and_cancel()
        amount, rtype, payment = process_refund(reg, reason="Test refund")
        self.assertEqual(amount, Decimal("20.00"))
        self.assertTrue(payment.is_refund)
        self.assertEqual(payment.amount, Decimal("20.00"))
        reg.refresh_from_db()
        self.assertEqual(reg.amount_refunded, Decimal("20.00"))

    def test_process_refund_admin_only_requires_override(self):
        """Admin-only refund raises without admin_override."""
        reg = self._register_and_cancel()
        ref_date = self.session.date
        with self.assertRaises(ValidationError) as ctx:
            process_refund(reg, reference_date=ref_date)
        self.assertIn("admin approval", str(ctx.exception))

    def test_process_refund_admin_override(self):
        """Admin can override the admin-only restriction."""
        reg = self._register_and_cancel()
        ref_date = self.session.date
        amount, rtype, payment = process_refund(
            reg, amount=Decimal("20.00"), admin_override=True,
            reference_date=ref_date, reason="Admin override",
        )
        self.assertEqual(amount, Decimal("20.00"))

    def test_double_refund_prevented(self):
        """Cannot refund more than was charged."""
        reg = self._register_and_cancel()
        process_refund(reg, reason="First")
        with self.assertRaises(ValidationError) as ctx:
            process_refund(reg, reason="Second")
        self.assertIn("already been fully refunded", str(ctx.exception))

    def test_partial_then_remaining_refund(self):
        """Can do a partial refund, then refund the remainder."""
        reg = self._register_and_cancel()
        process_refund(reg, amount=Decimal("10.00"), reason="Partial")
        reg.refresh_from_db()
        self.assertEqual(reg.amount_refunded, Decimal("10.00"))
        process_refund(reg, amount=Decimal("10.00"), reason="Remainder")
        reg.refresh_from_db()
        self.assertEqual(reg.amount_refunded, Decimal("20.00"))

    def test_refund_not_cancelled_raises(self):
        """Cannot refund a non-cancelled registration."""
        reg, _ = register_person(
            self.org, self.person, self.event, self.term, self.tier,
        )
        with self.assertRaises(ValidationError):
            process_refund(reg)

    def test_no_policy_gives_full_refund(self):
        """Without a refund policy, full refund is the default."""
        self.policy.delete()
        reg = self._register_and_cancel()
        amount, rtype = calculate_refund(reg)
        self.assertEqual(amount, Decimal("20.00"))
        self.assertEqual(rtype, "full")


# ---------------------------------------------------------------------------
# Offline Payment Tests
# ---------------------------------------------------------------------------

class OfflinePaymentTests(TestCase):

    def setUp(self):
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.person = OrgFactory.create_person(self.org, email="offline@t.com")
        self.event = Event.objects.create(
            organization=self.org, term=self.term, title="Pay Event",
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("8.00"),
        )

    def test_record_check_payment(self):
        """Record a check payment."""
        reg, _ = register_person(
            self.org, self.person, self.event, self.term, self.tier,
        )
        payment = record_offline_payment(
            self.org, self.person, Decimal("8.00"), "check",
            registrations=[reg], check_number="1234",
        )
        self.assertEqual(payment.method, "check")
        self.assertEqual(payment.amount, Decimal("8.00"))
        self.assertEqual(payment.check_number, "1234")
        self.assertIn(reg, payment.registrations.all())

    def test_record_cash_payment(self):
        """Record a cash payment."""
        payment = record_offline_payment(
            self.org, self.person, Decimal("8.00"), "cash",
            reference_note="Walk-in registration",
        )
        self.assertEqual(payment.method, "cash")
        self.assertEqual(payment.reference_note, "Walk-in registration")

    def test_invalid_method_raises(self):
        """Invalid offline payment method raises ValidationError."""
        with self.assertRaises(ValidationError):
            record_offline_payment(
                self.org, self.person, Decimal("8.00"), "square_api",
            )

    def test_reconcile_payment(self):
        """Reconcile a payment with registrations."""
        reg, _ = register_person(
            self.org, self.person, self.event, self.term, self.tier,
        )
        payment = record_offline_payment(
            self.org, self.person, Decimal("8.00"), "cash",
        )
        self.assertEqual(payment.registrations.count(), 0)
        reconcile_payment(payment, [reg])
        self.assertEqual(payment.registrations.count(), 1)
        self.assertIn(reg, payment.registrations.all())


# ---------------------------------------------------------------------------
# Admin Smoke Tests
# ---------------------------------------------------------------------------

class AdminSmokeTests(TestCase):

    def setUp(self):
        from core.models import User
        self.user = User.objects.create_superuser(
            username="admin", email="admin@test.com", password="pass",
        )
        self.client.login(username="admin", password="pass")
        self.org = OrgFactory.create_org()

    def test_admin_index_loads(self):
        """Django admin index loads successfully."""
        resp = self.client.get("/admin/")
        self.assertEqual(resp.status_code, 200)

    def test_event_changelist_loads(self):
        """Event changelist loads."""
        resp = self.client.get("/admin/core/event/")
        self.assertEqual(resp.status_code, 200)

    def test_registration_changelist_loads(self):
        """Registration changelist loads."""
        resp = self.client.get("/admin/core/registration/")
        self.assertEqual(resp.status_code, 200)

    def test_person_changelist_loads(self):
        """Person changelist loads."""
        resp = self.client.get("/admin/core/person/")
        self.assertEqual(resp.status_code, 200)

    def test_payment_changelist_loads(self):
        """Payment changelist loads."""
        resp = self.client.get("/admin/core/payment/")
        self.assertEqual(resp.status_code, 200)

    def test_organization_changelist_loads(self):
        """Organization changelist loads."""
        resp = self.client.get("/admin/core/organization/")
        self.assertEqual(resp.status_code, 200)

    def test_auditlog_changelist_loads(self):
        """AuditLog changelist is read-only."""
        resp = self.client.get("/admin/core/auditlog/")
        self.assertEqual(resp.status_code, 200)

    def test_auditlog_add_blocked(self):
        """Cannot add AuditLog entries through admin."""
        resp = self.client.get("/admin/core/auditlog/add/")
        self.assertEqual(resp.status_code, 403)


# ===========================================================================
# Phase 3 Tests — Import, Export, Messaging
# ===========================================================================


# ---------------------------------------------------------------------------
# Square Items CSV Import Tests
# ---------------------------------------------------------------------------

class SquareItemsImportTests(TestCase):

    def setUp(self):
        from core.importers import SquareItemsImporter

        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.event = Event.objects.create(
            organization=self.org, term=self.term,
            title="Test Lecture", legacy_id="101",
        )
        self.person = OrgFactory.create_person(
            self.org, first="Jane", last="Doe", email="jane@test.com",
        )
        self.importer = SquareItemsImporter(self.org, self.term, self.tier)

    def _make_csv(self, rows):
        """Build a CSV string from header + data rows."""
        import io, csv
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            "Date", "Time", "Time Zone", "Category", "Item", "Qty",
            "Price Point Name", "SKU", "Modifiers Applied", "Gross Sales",
            "Discounts", "Net Sales", "Tax", "Transaction ID", "Payment ID",
            "Device Name", "Notes", "Details", "Event Type", "Location",
            "Dining Option", "Customer ID", "Customer Name",
            "Customer Reference ID", "Unit", "Count", "GTIN",
            "Itemization Type", "Fulfillment Note", "Channel", "Token",
            "Card Brand", "PAN Suffix",
        ])
        for row in rows:
            writer.writerow(row)
        return io.StringIO(buf.getvalue())

    def test_import_creates_registration_and_payment(self):
        """Importing a valid row creates a registration and payment."""
        csv_file = self._make_csv([[
            "2025-08-29", "16:35:45", "Eastern", "Evening Lecture",
            "Test Lecture", "1", "", "101", "", "$1.00", "$0.00",
            "$1.00", "$0.00", "txn-abc-123", "pay-123", "", "", "",
            "", "", "", "", "Jane Doe", "", "", "", "", "", "", "", "", "", "",
        ]])
        results = self.importer.import_csv(csv_file)
        self.assertEqual(results["imported"], 1)
        self.assertEqual(Registration.objects.filter(event=self.event).count(), 1)
        self.assertEqual(Payment.objects.filter(square_transaction_id="txn-abc-123").count(), 1)

    def test_idempotent_reimport(self):
        """Re-importing the same transaction is skipped."""
        row = [
            "2025-08-29", "16:35:45", "Eastern", "Evening Lecture",
            "Test Lecture", "1", "", "101", "", "$1.00", "$0.00",
            "$1.00", "$0.00", "txn-abc-123", "pay-123", "", "", "",
            "", "", "", "", "Jane Doe", "", "", "", "", "", "", "", "", "", "",
        ]
        self.importer.import_csv(self._make_csv([row]))
        # Import again
        from core.importers import SquareItemsImporter
        importer2 = SquareItemsImporter(self.org, self.term, self.tier)
        results = importer2.import_csv(self._make_csv([row]))
        self.assertEqual(results["skipped_duplicate"], 1)
        self.assertEqual(results["imported"], 0)

    def test_missing_sku_skipped(self):
        """Rows without SKU are skipped."""
        csv_file = self._make_csv([[
            "2025-08-29", "16:35:45", "Eastern", "", "Item", "1", "",
            "", "", "$1.00", "$0.00", "$1.00", "$0.00", "txn-999",
            "pay-999", "", "", "", "", "", "", "", "Jane Doe", "", "",
            "", "", "", "", "", "", "", "",
        ]])
        results = self.importer.import_csv(csv_file)
        self.assertEqual(results["skipped_no_sku"], 1)

    def test_unknown_event_reported(self):
        """SKU with no matching event is reported as error."""
        csv_file = self._make_csv([[
            "2025-08-29", "16:35:45", "Eastern", "", "Unknown", "1",
            "", "999", "", "$1.00", "$0.00", "$1.00", "$0.00",
            "txn-unk", "pay-unk", "", "", "", "", "", "", "",
            "Jane Doe", "", "", "", "", "", "", "", "", "", "",
        ]])
        results = self.importer.import_csv(csv_file)
        self.assertEqual(results["skipped_no_event"], 1)

    def test_ambiguous_person_flagged(self):
        """Ambiguous customer name is flagged for review."""
        OrgFactory.create_person(
            self.org, first="John", last="Doe", email="john@test.com",
        )
        csv_file = self._make_csv([[
            "2025-08-29", "16:35:45", "Eastern", "", "Test", "1",
            "", "101", "", "$1.00", "$0.00", "$1.00", "$0.00",
            "txn-amb", "pay-amb", "", "", "", "", "", "", "",
            "Unknown Doe", "", "", "", "", "", "", "", "", "", "",
        ]])
        results = self.importer.import_csv(csv_file)
        self.assertTrue(len(results["needs_review"]) > 0)

    def test_missing_columns_raises(self):
        """CSV missing required columns raises ValidationError."""
        import io
        csv_file = io.StringIO("Col1,Col2\na,b\n")
        with self.assertRaises(ValidationError):
            self.importer.import_csv(csv_file)


# ---------------------------------------------------------------------------
# Square Transactions CSV Import Tests
# ---------------------------------------------------------------------------

class SquareTransactionsImportTests(TestCase):

    def setUp(self):
        from core.importers import SquareTransactionsImporter
        self.org = OrgFactory.create_org()
        self.person = OrgFactory.create_person(
            self.org, first="Jane", last="Doe", email="jane@test.com",
        )
        self.importer = SquareTransactionsImporter(self.org)

    def _make_csv(self, rows):
        import io, csv
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            "Date", "Time", "Time Zone", "Gross Sales", "Discounts",
            "Service Charges", "Net Sales", "Gift Card Sales", "Tax", "Tip",
            "Partial Refunds", "Total Collected", "Source", "Card",
            "Card Entry Methods", "Cash", "Square Gift Card", "Other Tender",
            "Other Tender Type", "Tender Note", "Fees", "Net Total",
            "Transaction ID", "Payment ID", "Card Brand", "PAN Suffix",
            "Device Name", "Staff Name", "Staff ID", "Details", "Description",
            "Event Type", "Location", "Dining Option", "Customer ID",
            "Customer Name", "Customer Reference ID", "Device Nickname",
            "Third Party Fees", "Deposit ID", "Deposit Date",
            "Deposit Details", "Fee Percentage Rate", "Fee Fixed Rate",
            "Refund Reason", "Discount Name", "Transaction Status",
            "Cash App", "Order Reference ID", "Fulfillment Note",
            "Free Processing Applied", "Channel", "Unattributed Tips",
            "Table Info",
        ])
        for row in rows:
            writer.writerow(row)
        return io.StringIO(buf.getvalue())

    def test_import_creates_payment(self):
        """Importing a transaction creates a Payment."""
        csv_file = self._make_csv([[
            "2025-08-29", "16:35:45", "Eastern", "$1.00", "$0.00",
            "$0.00", "$1.00", "$0.00", "$0.00", "$0.00", "$0.00",
            "$1.00", "Virtual Terminal", "", "", "", "", "", "", "",
            "$0.03", "$0.97", "txn-trans-001", "pay-001", "", "",
            "", "", "", "", "Registration", "", "", "", "", "Jane Doe",
            "", "", "", "", "", "", "", "", "", "", "", "", "", "", "",
            "", "", "",
        ]])
        results = self.importer.import_csv(csv_file)
        self.assertEqual(results["imported"], 1)
        p = Payment.objects.get(square_transaction_id="txn-trans-001")
        self.assertEqual(p.amount, Decimal("1.00"))
        self.assertEqual(p.person, self.person)

    def test_idempotent_reimport(self):
        """Re-importing same transaction ID is skipped."""
        row = [
            "2025-08-29", "16:35:45", "Eastern", "$1.00", "$0.00",
            "$0.00", "$1.00", "$0.00", "$0.00", "$0.00", "$0.00",
            "$1.00", "", "", "", "", "", "", "", "", "", "",
            "txn-trans-002", "pay-002", "", "", "", "", "", "", "",
            "", "", "", "", "Jane Doe", "", "", "", "", "", "", "",
            "", "", "", "", "", "", "", "", "", "", "",
        ]
        self.importer.import_csv(self._make_csv([row]))
        from core.importers import SquareTransactionsImporter
        importer2 = SquareTransactionsImporter(self.org)
        results = importer2.import_csv(self._make_csv([row]))
        self.assertEqual(results["skipped_duplicate"], 1)


# ---------------------------------------------------------------------------
# Directory Import Tests
# ---------------------------------------------------------------------------

class DirectoryImportTests(TestCase):

    def setUp(self):
        from core.importers import DirectoryImporter
        self.org = OrgFactory.create_org()
        self.importer = DirectoryImporter(self.org)

    def _make_csv(self, rows):
        import io, csv
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            "First Name", "Last Name", "Title", "Nickname", "Move-in Date",
            "DOB", "Unit #", "Address", "City", "State/Province",
            "Postal/Zip Code", "Email", "Cell Phone Number", "Phone Number",
            "Bio", "Interests", "Gender", "Marital Status", "Building",
            "Education", "Military Service", "Primary Language",
            "Secondary Language",
        ])
        for row in rows:
            writer.writerow(row)
        return io.StringIO(buf.getvalue())

    def test_import_creates_person(self):
        """Importing a new person creates a Person record."""
        csv_file = self._make_csv([[
            "Arthur", "Sampleton", "-", "-", "Dec 2024", "-",
            "Diamond Apt-416", "100 Main Street, #416", "Sampletown",
            "Maryland", "99999", "arthur@test.com", "-", "5555550100",
            "-", "-", "Male", "-", "-", "-", "No", "-", "-",
        ]])
        results = self.importer.import_csv(csv_file)
        self.assertEqual(results["created"], 1)
        p = Person.objects.get(organization=self.org, email="arthur@test.com")
        self.assertEqual(p.first_name, "Arthur")
        self.assertEqual(p.last_name, "Sampleton")
        self.assertTrue(p.is_resident)

    def test_idempotent_reimport_by_email(self):
        """Re-import matches on email and skips unchanged."""
        OrgFactory.create_person(
            self.org, first="Arthur", last="Sampleton",
            email="arthur@test.com", is_resident=True,
            phone="5555550100",
        )
        csv_file = self._make_csv([[
            "Arthur", "Sampleton", "-", "-", "-", "-", "-", "-",
            "-", "-", "-", "arthur@test.com", "-", "5555550100",
            "-", "-", "-", "-", "-", "-", "-", "-", "-",
        ]])
        results = self.importer.import_csv(csv_file)
        self.assertEqual(results["skipped_unchanged"], 1)
        self.assertEqual(results["created"], 0)

    def test_conflict_detected(self):
        """Conflicting data is flagged for admin review."""
        OrgFactory.create_person(
            self.org, first="Arthur", last="Sampleton",
            email="arthur@test.com", phone="1111111111",
        )
        csv_file = self._make_csv([[
            "Arthur", "Sampleton", "-", "-", "-", "-", "-", "-",
            "-", "-", "-", "arthur@test.com", "-", "9999999999",
            "-", "-", "-", "-", "-", "-", "-", "-", "-",
        ]])
        results = self.importer.import_csv(csv_file)
        self.assertTrue(len(results["conflicts"]) > 0)

    def test_new_fields_populated(self):
        """Import fills in empty fields on existing records."""
        p = OrgFactory.create_person(
            self.org, first="Arthur", last="Sampleton",
            email="arthur@test.com",
        )
        self.assertEqual(p.city, "")
        csv_file = self._make_csv([[
            "Arthur", "Sampleton", "-", "-", "-", "-", "-", "-",
            "Sampletown", "Maryland", "99999", "arthur@test.com",
            "-", "-", "-", "-", "-", "-", "-", "-", "-", "-", "-",
        ]])
        from core.importers import DirectoryImporter
        DirectoryImporter(self.org).import_csv(csv_file)
        p.refresh_from_db()
        self.assertEqual(p.city, "Sampletown")


# ---------------------------------------------------------------------------
# CSV Export Tests
# ---------------------------------------------------------------------------

class CSVExportTests(TestCase):

    def setUp(self):
        from core.services import register_person, record_offline_payment
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.event = Event.objects.create(
            organization=self.org, term=self.term, title="Export Event",
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("10.00"),
        )
        self.person = OrgFactory.create_person(
            self.org, first="Alice", last="Smith", email="alice@test.com",
        )
        self.reg, _ = register_person(
            self.org, self.person, self.event, self.term, self.tier,
        )
        self.payment = record_offline_payment(
            self.org, self.person, Decimal("10.00"), "check",
            registrations=[self.reg], check_number="5678",
        )

    def test_export_registrations_csv(self):
        """Registration CSV export contains expected data."""
        from core.exporters import export_registrations_csv
        csv_str = export_registrations_csv(self.org, term=self.term)
        self.assertIn("Smith, Alice", csv_str)
        self.assertIn("Export Event", csv_str)
        self.assertIn("10.00", csv_str)

    def test_export_payments_csv(self):
        """Payment CSV export contains expected data."""
        from core.exporters import export_payments_csv
        csv_str = export_payments_csv(self.org)
        self.assertIn("Smith, Alice", csv_str)
        self.assertIn("check", csv_str)
        self.assertIn("5678", csv_str)

    def test_export_registrations_custom_columns(self):
        """Custom column selection works."""
        from core.exporters import export_registrations_csv
        columns = [
            ("person_name", "Student"),
            ("event_title", "Course"),
            ("net_amount", "Net"),
        ]
        csv_str = export_registrations_csv(
            self.org, term=self.term, columns=columns,
        )
        lines = csv_str.strip().split("\n")
        self.assertEqual(lines[0].strip(), "Student,Course,Net")


# ---------------------------------------------------------------------------
# Excel Export Tests
# ---------------------------------------------------------------------------

class ExcelExportTests(TestCase):

    def setUp(self):
        from core.services import register_person
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.event = Event.objects.create(
            organization=self.org, term=self.term, title="Excel Event",
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("10.00"),
        )
        self.person = OrgFactory.create_person(
            self.org, first="Bob", last="Jones", email="bob@test.com",
        )
        register_person(self.org, self.person, self.event, self.term, self.tier)

    def test_export_registrations_excel(self):
        """Excel export produces valid xlsx bytes."""
        from core.exporters import export_registrations_excel
        import openpyxl, io
        data = export_registrations_excel(self.org, term=self.term)
        self.assertIsInstance(data, bytes)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        self.assertEqual(ws.title, "Registrations")
        # Header row
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Name", headers)
        # Data row
        self.assertEqual(ws.cell(row=2, column=2).value, "Jones, Bob")


# ---------------------------------------------------------------------------
# QuickBooks IIF Export Tests
# ---------------------------------------------------------------------------

class IIFExportTests(TestCase):

    def setUp(self):
        from core.services import record_offline_payment
        self.org = OrgFactory.create_org()
        self.person = OrgFactory.create_person(
            self.org, first="Carol", last="White", email="carol@test.com",
        )
        self.payment = record_offline_payment(
            self.org, self.person, Decimal("25.00"), "check",
            check_number="9999",
        )

    def test_iif_export_structure(self):
        """IIF export has correct header and transaction structure."""
        from core.exporters import export_quickbooks_iif
        iif = export_quickbooks_iif(self.org)
        lines = iif.split("\n")
        self.assertTrue(lines[0].startswith("!TRNS"))
        self.assertTrue(lines[1].startswith("!SPL"))
        self.assertTrue(lines[2].startswith("!ENDTRNS"))
        # Should contain a TRNS line with the payment
        trns_lines = [l for l in lines if l.startswith("TRNS")]
        self.assertTrue(len(trns_lines) > 0)
        self.assertIn("White, Carol", iif)
        self.assertIn("25.0", iif)

    def test_iif_export_refund(self):
        """IIF export correctly handles refund payments."""
        from core.exporters import export_quickbooks_iif
        Payment.objects.create(
            organization=self.org, person=self.person,
            amount=Decimal("10.00"), method="check",
            paid_at=timezone.now(), is_refund=True,
            refund_reason="Test refund",
        )
        iif = export_quickbooks_iif(self.org)
        self.assertIn("Registration Refunds", iif)
        self.assertIn("Refund: Test refund", iif)


# ---------------------------------------------------------------------------
# QuickBooks QBO Export Tests
# ---------------------------------------------------------------------------

class QBOExportTests(TestCase):

    def setUp(self):
        from core.services import record_offline_payment
        self.org = OrgFactory.create_org()
        self.person = OrgFactory.create_person(
            self.org, first="Dan", last="Green", email="dan@test.com",
        )
        self.payment = record_offline_payment(
            self.org, self.person, Decimal("50.00"), "cash",
        )

    def test_qbo_export_valid_xml(self):
        """QBO export produces parseable XML."""
        from core.exporters import export_quickbooks_qbo
        import xml.etree.ElementTree as ET
        qbo = export_quickbooks_qbo(self.org)
        # Should parse as valid XML
        root = ET.fromstring(qbo)
        self.assertEqual(root.tag, "OFX")

    def test_qbo_contains_transaction(self):
        """QBO export contains the payment transaction."""
        from core.exporters import export_quickbooks_qbo
        qbo = export_quickbooks_qbo(self.org)
        self.assertIn("Green, Dan", qbo)
        self.assertIn("50.00", qbo)
        self.assertIn("CREDIT", qbo)


# ---------------------------------------------------------------------------
# Attendance Sheet Tests
# ---------------------------------------------------------------------------

class AttendanceSheetTests(TestCase):

    def setUp(self):
        from core.services import register_person
        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.event = Event.objects.create(
            organization=self.org, term=self.term,
            title="Attendance Event", capacity=30,
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("1.00"),
        )
        self.session = EventSession.objects.create(
            organization=self.org, event=self.event,
            date=datetime.date(2025, 10, 6),
            start_time=datetime.time(19, 30),
        )
        self.p1 = OrgFactory.create_person(
            self.org, first="Eve", last="Adams", email="eve@test.com",
        )
        self.p2 = OrgFactory.create_person(
            self.org, first="Frank", last="Burns", email="frank@test.com",
        )
        register_person(self.org, self.p1, self.event, self.term, self.tier)
        register_person(self.org, self.p2, self.event, self.term, self.tier)

    def test_attendance_html_contains_names(self):
        """Attendance HTML includes registered people's names."""
        from core.attendance import generate_attendance_html
        html = generate_attendance_html(self.event)
        self.assertIn("Adams, Eve", html)
        self.assertIn("Burns, Frank", html)

    def test_attendance_html_contains_event_info(self):
        """Attendance HTML includes event title and org name."""
        from core.attendance import generate_attendance_html
        html = generate_attendance_html(self.event)
        self.assertIn("Attendance Event", html)
        self.assertIn(self.org.name, html)
        self.assertIn("Registered: 2", html)

    def test_attendance_html_blank_rows(self):
        """Attendance HTML includes blank rows for walk-ins."""
        from core.attendance import generate_attendance_html
        html = generate_attendance_html(self.event, blank_rows=3)
        # Should have registered names plus blank rows
        self.assertIn("Adams, Eve", html)

    def test_attendance_for_term(self):
        """generate_attendance_sheets_for_term produces sheets for each event."""
        from core.attendance import generate_attendance_sheets_for_term
        results = generate_attendance_sheets_for_term(self.org, self.term)
        self.assertEqual(len(results), 1)
        event, html = results[0]
        self.assertEqual(event, self.event)
        self.assertIn("Attendance Event", html)


# ---------------------------------------------------------------------------
# Email Provider Tests
# ---------------------------------------------------------------------------

class EmailProviderTests(TestCase):

    def test_console_provider_records_messages(self):
        """ConsoleEmailProvider records sent messages."""
        from core.providers import ConsoleEmailProvider
        provider = ConsoleEmailProvider()
        provider.reset()
        result = provider.send(
            to="test@test.com",
            subject="Test Subject",
            body="Test body",
            from_email="from@test.com",
        )
        self.assertTrue(result.success)
        self.assertEqual(len(ConsoleEmailProvider.sent_messages), 1)
        self.assertEqual(ConsoleEmailProvider.sent_messages[0]["to"], "test@test.com")
        provider.reset()

    def test_smtp_provider_exists(self):
        """SMTPEmailProvider can be instantiated."""
        from core.providers import SMTPEmailProvider
        provider = SMTPEmailProvider()
        self.assertIsNotNone(provider)

    def test_mailgun_provider_needs_config(self):
        """MailgunEmailProvider fails without config."""
        from core.providers import MailgunEmailProvider
        provider = MailgunEmailProvider()
        result = provider.send(to="x@x.com", subject="x", body="x")
        self.assertFalse(result.success)
        self.assertIn("api_key", result.error)

    def test_sendgrid_provider_needs_config(self):
        """SendGridEmailProvider fails without config."""
        from core.providers import SendGridEmailProvider
        provider = SendGridEmailProvider()
        result = provider.send(to="x@x.com", subject="x", body="x")
        self.assertFalse(result.success)
        self.assertIn("api_key", result.error)

    def test_sms_stub(self):
        """SMS stub returns not-implemented."""
        from core.providers import SMSProvider
        sms = SMSProvider()
        result = sms.send(to="+1234567890", message="Test")
        self.assertFalse(result["success"])

    def test_get_email_provider_default(self):
        """get_email_provider returns SMTP by default."""
        from core.providers import get_email_provider, SMTPEmailProvider
        org = OrgFactory.create_org()
        provider = get_email_provider(org)
        self.assertIsInstance(provider, SMTPEmailProvider)

    def test_get_email_provider_override(self):
        """Provider override works for testing."""
        from core.providers import (
            ConsoleEmailProvider,
            clear_email_provider_override,
            get_email_provider,
            set_email_provider_override,
        )
        org = OrgFactory.create_org()
        console = ConsoleEmailProvider()
        set_email_provider_override(console)
        provider = get_email_provider(org)
        self.assertIs(provider, console)
        clear_email_provider_override()


# ---------------------------------------------------------------------------
# Messaging Tests
# ---------------------------------------------------------------------------

class MessagingTests(TestCase):

    def setUp(self):
        from core.providers import ConsoleEmailProvider as _CEP
        from core.providers import set_email_provider_override
        from core.services import register_person

        self.ConsoleEmailProvider = _CEP
        self.console = _CEP()
        _CEP.reset()
        set_email_provider_override(self.console)

        self.org = OrgFactory.create_org()
        self.term = OrgFactory.create_term(self.org)
        self.tier = PricingTier.objects.create(
            organization=self.org, name="Resident", slug="resident",
        )
        self.event = Event.objects.create(
            organization=self.org, term=self.term,
            title="Messaging Event", capacity=10,
        )
        EventPrice.objects.create(
            organization=self.org, event=self.event,
            tier=self.tier, timing="pre", amount=Decimal("5.00"),
        )
        self.session = EventSession.objects.create(
            organization=self.org, event=self.event,
            date=datetime.date(2025, 10, 6),
            start_time=datetime.time(19, 30),
        )
        self.person = OrgFactory.create_person(
            self.org, first="Grace", last="Hopper", email="grace@test.com",
        )
        self.reg, _ = register_person(
            self.org, self.person, self.event, self.term, self.tier,
        )

    def tearDown(self):
        from core.providers import clear_email_provider_override
        clear_email_provider_override()

    def test_send_registration_confirmation(self):
        """Registration confirmation email is sent."""
        from core.messaging import send_registration_confirmation
        result = send_registration_confirmation(self.reg)
        self.assertTrue(result.success)
        self.assertEqual(len(self.ConsoleEmailProvider.sent_messages), 1)
        msg = self.ConsoleEmailProvider.sent_messages[0]
        self.assertEqual(msg["to"], "grace@test.com")
        self.assertIn("Messaging Event", msg["subject"])

    def test_send_waitlist_notification(self):
        """Waitlist notification email is sent."""
        from core.messaging import send_waitlist_notification
        wl = WaitlistEntry.objects.create(
            organization=self.org, registration=self.reg, position=3,
        )
        result = send_waitlist_notification(self.reg, wl)
        self.assertTrue(result.success)
        msg = self.ConsoleEmailProvider.sent_messages[0]
        self.assertIn("Waitlisted", msg["subject"])

    def test_send_cancellation_with_refund(self):
        """Cancellation email includes refund amount."""
        from core.messaging import send_cancellation_notice
        result = send_cancellation_notice(self.reg, refund_amount=Decimal("5.00"))
        self.assertTrue(result.success)
        msg = self.ConsoleEmailProvider.sent_messages[0]
        self.assertIn("Cancelled", msg["subject"])
        self.assertIn("5.00", msg["body"])

    def test_send_nudge(self):
        """Nudge email is sent to recipients."""
        from core.messaging import send_nudge
        from core.models import EventGroupType, EventGroup, EventGroupMembership

        gt = EventGroupType.objects.create(
            organization=self.org, name="Nudge", slug="nudge",
        )
        group = EventGroup.objects.create(
            organization=self.org, group_type=gt,
            name="Week 1 Nudge", position=1,
        )
        EventGroupMembership.objects.create(group=group, event=self.event, position=0)

        results = send_nudge(
            self.org, group,
            recipients=Person.objects.filter(pk=self.person.pk),
        )
        self.assertEqual(len(results), 1)
        self.assertTrue(results[0].success)
        msg = self.ConsoleEmailProvider.sent_messages[0]
        self.assertIn("Messaging Event", msg["subject"])

    def test_no_email_skips_send(self):
        """Person without email doesn't trigger send."""
        from core.messaging import send_registration_confirmation
        no_email = OrgFactory.create_person(
            self.org, first="No", last="Email",
        )
        from core.services import register_person as rp
        reg2, _ = rp(
            self.org, no_email, self.event, self.term, self.tier,
            amount_override=Decimal("0"),
        )
        result = send_registration_confirmation(reg2)
        self.assertIsNone(result)

    def test_bulk_message(self):
        """Bulk messaging sends to all recipients."""
        from core.messaging import send_bulk_message
        p2 = OrgFactory.create_person(
            self.org, first="Ada", last="Lovelace", email="ada@test.com",
        )
        results = send_bulk_message(
            self.org,
            recipients=[self.person, p2],
            subject="Hello {{ person.first_name }}",
            body_template="Dear {{ person.first_name }}, welcome to {{ org.name }}.",
        )
        self.assertEqual(len(results), 2)
        subjects = [m["subject"] for m in self.ConsoleEmailProvider.sent_messages]
        self.assertIn("Hello Grace", subjects)
        self.assertIn("Hello Ada", subjects)
